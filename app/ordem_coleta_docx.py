"""Preenche o modelo .docx real da Ordem de Coleta (AMG Transportes / AMGLOG
/ JVA, conforme a empresa/filial escolhida) com os dados que o sistema tem.

Campos operacionais que so a Logistica sabe depois de despachar (motorista,
placas, ANTT, CT-e, adiantamento/saldo, liberacao Buonny) ficam em branco no
documento, prontos para preenchimento manual -- nao inventamos esses dados.
"""

from __future__ import annotations

import copy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.table import Table, _Cell, _Row

from app.constants import COMERCIAL_RESPONSAVEL
from app.filiais import EMPRESA_TEMPLATE
from app.models import Quote
from app.utils import format_brl

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "assets" / "ordem_coleta"

LABEL_COLETA = "ENDEREÇOS – LOCAL DA COLETA"
LABEL_ENTREGA = "ENDEREÇOS – LOCAL DA ENTREGA"
LABEL_DADOS_CARGA = "DADOS DA CARGA"


def _endereco(cidade: str | None, cep, endereco, numero, bairro) -> str:
    endereco_completo = f"{endereco}, {numero}" if endereco and numero else endereco
    partes = [p for p in [endereco_completo, bairro, cidade, f"CEP {cep}" if cep else None] if p]
    return ", ".join(partes) if partes else "-"


def _first_value_cell(row: _Row) -> _Cell | None:
    """Primeira celula, andando da esquerda, cujo bloco de merge e diferente
    do rotulo (cells[0]). Funciona tanto para linhas de campo unico (o valor
    e um merge unico do resto da linha) quanto para linhas com dois campos
    lado a lado (o valor do 1o campo fica logo apos o rotulo, antes do 2o
    rotulo comecar)."""
    label_tc = row.cells[0]._tc
    for cell in row.cells[1:]:
        if cell._tc is not label_tc:
            return cell
    return None


def _last_value_cell(row: _Row) -> _Cell | None:
    """Ultima celula da linha, se for distinta do rotulo -- usado para o 2o
    campo ("N. DA COTACAO/COLETA") que divide a linha com "COMERCIAL
    RESPONSAVEL"."""
    target = row.cells[-1]
    return target if target._tc is not row.cells[0]._tc else None


def _set_first(row: _Row, texto) -> None:
    target = _first_value_cell(row)
    if target is not None:
        target.text = str(texto) if texto not in (None, "") else "-"


def _set_last(row: _Row, texto) -> None:
    target = _last_value_cell(row)
    if target is not None:
        target.text = str(texto) if texto not in (None, "") else "-"


def _indice_por_label(table: Table, label_exato: str) -> int | None:
    for i, row in enumerate(table.rows):
        if row.cells[0].text.strip() == label_exato:
            return i
    return None


def _clonar_bloco(table: Table, inicio: int, fim: int, vezes_extra: int) -> None:
    """Duplica table.rows[inicio:fim] (cabecalho + linhas de dado + linha em
    branco) `vezes_extra` vezes, inserindo cada copia logo apos a anterior.
    Mexe direto no XML da tabela (copy.deepcopy do elemento <w:tr> + addnext)
    -- python-docx nao tem um jeito de "duplicar linha" pronto, mas <w:tr>
    nao exige id unico no OOXML, entao clonar o elemento inteiro (preserva
    merge de celulas e formatacao) e valido. Depois de chamada, table.rows
    ja reflete as linhas novas (e' recalculado do XML a cada acesso)."""
    originais = [row._tr for row in table.rows[inicio:fim]]
    ultimo = originais[-1]
    for _ in range(vezes_extra):
        anterior = ultimo
        for tr in originais:
            clone = copy.deepcopy(tr)
            anterior.addnext(clone)
            anterior = clone
        ultimo = anterior


def _numerar_blocos(table: Table, label_base: str) -> None:
    """Depois de clonado, todo bloco (original + copias) ainda tem o mesmo
    rotulo sem numero -- numera em ordem (' 1', ' 2', ...) so quando ha mais
    de 1, pra cotacao com uma coleta/um destino so (o caso comum) manter o
    rotulo original do modelo, sem numero."""
    n = 0
    for row in table.rows:
        if row.cells[0].text.strip() == label_base:
            n += 1
            row.cells[0].text = f"{label_base} {n}"


def _expandir_locais(doc: Document, quote: Quote) -> None:
    """Replica o bloco 'LOCAL DA COLETA'/'LOCAL DA ENTREGA' do modelo, uma
    vez pra cada coleta/destino alem do principal (ate 5 de cada, ver
    app.constants.MAX_LOCAIS_POR_TIPO) -- o modelo .docx fisico so tem 1
    bloco de cada por padrao; isso o repete dinamicamente a partir do MESMO
    arquivo-base, sem precisar de um template por combinacao possivel."""
    total_coletas = 1 + len(quote.coletas_extras)
    total_destinos = 1 + len(quote.destinos_extras)
    if total_coletas == 1 and total_destinos == 1:
        return  # caso comum -- nada a expandir

    for table in doc.tables:
        idx_coleta = _indice_por_label(table, LABEL_COLETA)
        if idx_coleta is None:
            continue  # tabela sem os campos de endereco -- nao e essa

        if total_coletas > 1:
            idx_entrega = _indice_por_label(table, LABEL_ENTREGA)
            _clonar_bloco(table, idx_coleta, idx_entrega, total_coletas - 1)

        if total_destinos > 1:
            # reconsulta: inserir as coletas extras empurrou o indice da
            # entrega pra baixo.
            idx_entrega = _indice_por_label(table, LABEL_ENTREGA)
            idx_dados_carga = _indice_por_label(table, LABEL_DADOS_CARGA)
            if idx_dados_carga is not None:
                _clonar_bloco(table, idx_entrega, idx_dados_carga, total_destinos - 1)

        if total_coletas > 1:
            _numerar_blocos(table, LABEL_COLETA)
        if total_destinos > 1:
            _numerar_blocos(table, LABEL_ENTREGA)
        break  # achou a tabela certa, nao precisa olhar as outras


def build_ordem_coleta_docx(quote: Quote) -> bytes:
    oc = quote.ordem_coleta
    solicitacao = quote.solicitacao
    proposal = quote.proposal
    template_name = EMPRESA_TEMPLATE.get(oc.empresa, "amg_transportes.docx")
    doc = Document(str(TEMPLATES_DIR / template_name))
    _expandir_locais(doc, quote)

    numero_cotacao = f"{quote.code} / {oc.numero}"
    campos_simples = {
        "N° DA COTAÇÃO/COLETA:": numero_cotacao,
        "DATA E HORA DA COLETA": quote.data_coleta.strftime("%d/%m/%Y"),
        "DATA E HORA DE ENTREGA": (
            quote.data_entrega.strftime("%d/%m/%Y") if quote.data_entrega else "-"
        ),
        "TIPO DE VEÍCULO": quote.tipo_veiculo or "-",
        "TIPO DO MATERIAL": quote.tipo_material
        + (f" - {quote.descricao_material}" if quote.descricao_material else ""),
        "VOLUME DO MATERIAL": str(quote.qtd_volumes) if quote.qtd_volumes else "-",
        "EMAIL PARA ENVIO BOLETO": quote.client_email,
        "TELEFONE DO SOLICITANTE": quote.client_phone or "-",
        "SOLICITANTE DO FRETE (NOME)": quote.client_name,
        "CNPJ PAGADOR DO FRETE": solicitacao.pagador_documento if solicitacao else "-",
        "FE TOTAL EMPRESA R$:": format_brl(proposal.valor_final) if proposal else "-",
    }

    for table in doc.tables:
        secao: str | None = None
        coleta_idx = 0  # qual bloco de coleta (1 = principal, 2+ = extras)
        entrega_idx = 0  # qual bloco de entrega (1 = principal, 2+ = extras)
        for row in table.rows:
            label = row.cells[0].text.strip()
            if not label:
                continue

            if label.startswith("ENDEREÇOS"):
                if "COLETA" in label:
                    secao = "coleta"
                    coleta_idx += 1
                elif "ENTREGA" in label:
                    secao = "entrega"
                    entrega_idx += 1
                else:
                    secao = None
                continue

            if secao == "coleta" and label == "NOME DA EMPRESA":
                # mesma empresa do cliente em qualquer coleta extra -- so o
                # endereco fisico muda de bloco pra bloco, nao quem pede.
                _set_first(row, quote.client_company)
            elif secao == "entrega" and label == "NOME DA EMPRESA":
                # idem: o destinatario informado na Solicitacao de Frete e
                # unico, vale pra qualquer entrega extra.
                _set_first(row, solicitacao.destinatario_nome if solicitacao else "-")
            elif secao == "coleta" and "ENDEREÇO COMPLETO" in label:
                if coleta_idx == 1:
                    _set_first(row, _endereco(
                        quote.origem_cidade, quote.origem_cep, quote.origem_endereco,
                        quote.origem_numero, quote.origem_bairro,
                    ))
                else:
                    loc = quote.coletas_extras[coleta_idx - 2]  # extras comecam no ordem=2
                    _set_first(row, _endereco(loc.cidade, loc.cep, loc.endereco, loc.numero, loc.bairro))
            elif secao == "entrega" and "ENDEREÇO COMPLETO" in label:
                if entrega_idx == 1:
                    _set_first(row, _endereco(
                        quote.destino_cidade, quote.destino_cep, quote.destino_endereco,
                        quote.destino_numero, quote.destino_bairro,
                    ))
                else:
                    loc = quote.destinos_extras[entrega_idx - 2]
                    _set_first(row, _endereco(loc.cidade, loc.cep, loc.endereco, loc.numero, loc.bairro))
            elif secao == "coleta" and label == "CNPJ COLETA":
                # CNPJ/CPF do CLIENTE (quem esta na ponta da coleta) -- nao
                # confundir com oc.cnpj_filial, que e o CNPJ da propria AMG
                # (usado so pra escolher o modelo/letterhead do documento).
                # Mesmo CNPJ em qualquer coleta extra, pelo mesmo motivo do
                # "NOME DA EMPRESA" acima.
                _set_first(row, quote.client_cnpj)
            elif label == "COMERCIAL RESPONSÁVEL":
                # a mesma linha carrega o rotulo "N. DA COTACAO/COLETA" mais
                # a frente, com seu proprio valor no ultimo slot da linha.
                # Comercial responsavel e sempre o mesmo, fixo (nao varia
                # por cotacao/filial).
                _set_first(row, COMERCIAL_RESPONSAVEL)
                _set_last(row, numero_cotacao)
            elif label == "PESO":
                _set_first(row, quote.peso_total)
            elif label in campos_simples:
                _set_first(row, campos_simples[label])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
