"""Testes de app/ordem_coleta_docx.py: preenchimento do modelo .docx real da
Ordem de Coleta a partir dos dados da cotacao."""

from datetime import date, timedelta

from docx import Document

from conftest import csrf_from

VALID = {
    "tipo_cotacao": "completa",
    "client_name": "Joao Alves",
    "client_company": "Construtora Alfa",
    "client_cnpj": "11.222.333/0001-44",
    "client_email": "compras@alfa.com.br",
    "client_phone": "(41) 99999-0000",
    "origem_cidade": "Curitiba - PR",
    "origem_cep": "80000-000",
    "origem_endereco": "Rua das Flores",
    "origem_numero": "100",
    "origem_bairro": "Centro",
    "destino_cidade": "Londrina - PR",
    "destino_cep": "86000-000",
    "destino_endereco": "Av. Brasil",
    "destino_numero": "200",
    "destino_bairro": "Jardim",
    "tipo_material": "Andaime",
    "qtd_volumes": "10",
    "peso_total": "5.000 kg",
    "valor_nf": "25.000,00",
    "tipo_veiculo": "Carreta Sider",
    "carroceria": "Sider",
    "capacidade_aprox": "28 toneladas",
    "servico_diaria": "sim",
    "servico_guincho": "nao",
}

PROPOSTA = {
    "custo_motorista": "8.000,00", "custo_pedagio": "900,00", "custo_impostos": "700,00",
    "custo_outros_internos": "0,00", "margem_pct": "20",
    "custos_adicionais": "0,00", "custos_adicionais_desc": "",
    "prazo_entrega": "1 dia útil", "observacoes": "",
}
# FC = 9600 ; FE = 9600 * 1.20 = 11520,00


def _fluxo_ate_oc(client, **proposta_overrides):
    page = client.get("/cotacao/nova?tipo=completa")
    token = csrf_from(page.text)
    payload = dict(VALID)
    payload["csrf_token"] = token
    payload["data_coleta"] = (date.today() + timedelta(days=5)).isoformat()
    resp = client.post("/cotacao/nova", data=payload, follow_redirects=False)
    code = resp.headers["location"].split("/cotacao/")[1].split("/")[0]

    admin_page = client.get("/admin/login")
    token = csrf_from(admin_page.text)
    client.post(
        "/admin/login",
        data={"username": "admin@amg.test", "password": "senha-teste",
              "csrf_token": token, "next": "/admin"},
        follow_redirects=False,
    )

    resp_page = client.get(f"/admin/cotacao/{code}")
    data = dict(PROPOSTA)
    data.update(proposta_overrides)
    data["csrf_token"] = csrf_from(resp_page.text)
    data["validade"] = (date.today() + timedelta(days=7)).isoformat()
    client.post(f"/admin/cotacao/{code}", data=data, follow_redirects=False)

    page = client.get(f"/admin/cotacao/{code}")
    token = csrf_from(page.text)
    client.post(f"/admin/cotacao/{code}/aprovar", data={"csrf_token": token}, follow_redirects=False)

    form_page = client.get(f"/admin/cotacao/{code}/gerar-oc")
    token = csrf_from(form_page.text)
    from app.filiais import extrair_uf, sugestao_empresa_cnpj

    empresa_cnpj = sugestao_empresa_cnpj(extrair_uf(VALID["destino_cidade"]))
    resp = client.post(
        f"/admin/cotacao/{code}/gerar-oc",
        data={
            "csrf_token": token, "empresa_cnpj": empresa_cnpj,
            # o form real vem pre-preenchido com o endereco da propria
            # cotacao (confirmavel/editavel); o navegador reenvia tudo --
            # simula isso aqui em vez de so mandar as cidades.
            "origem_cidade": VALID["origem_cidade"], "origem_cep": VALID["origem_cep"],
            "origem_endereco": VALID["origem_endereco"], "origem_numero": VALID["origem_numero"],
            "origem_bairro": VALID["origem_bairro"],
            "destino_cidade": VALID["destino_cidade"], "destino_cep": VALID["destino_cep"],
            "destino_endereco": VALID["destino_endereco"], "destino_numero": VALID["destino_numero"],
            "destino_bairro": VALID["destino_bairro"],
            "pagador": "Remetente", "pagador_documento": "12.345.678/0001-90",
            "destinatario_nome": "Obra Y", "destinatario_cnpj": "98.765.432/0001-10",
            "valor_nf": "25.000,00",
        },
        follow_redirects=False,
    )
    assert resp.status_code == 303, resp.text
    return code


def _tabelas_texto(doc: Document) -> str:
    partes = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                partes.append(cell.text)
    return "\n".join(partes)


def _linhas_ordenadas(doc: Document) -> list[tuple[str, str | None]]:
    """[(rotulo, valor), ...] da 1a tabela, na ordem do documento -- mesma
    logica de leitura de _set_first (1a celula distinta do rotulo). Lista
    (nao dict) de proposito: rotulos como "NOME DA EMPRESA" se repetem a
    cada bloco de coleta/entrega clonado, um dict perderia as repeticoes."""
    linhas = []
    table = doc.tables[0]
    for row in table.rows:
        label = row.cells[0].text.strip()
        if not label:
            continue
        label_tc = row.cells[0]._tc
        valor = None
        for cell in row.cells[1:]:
            if cell._tc is not label_tc:
                valor = cell.text.strip()
                break
        linhas.append((label, valor))
    return linhas


def test_build_ordem_coleta_docx_preenche_campos_da_cotacao(client):
    from app.crud import get_quote_by_code
    from app.database import SessionLocal
    from app.ordem_coleta_docx import build_ordem_coleta_docx

    code = _fluxo_ate_oc(client)
    with SessionLocal() as db:
        quote = get_quote_by_code(db, code)
        conteudo = build_ordem_coleta_docx(quote)

    from io import BytesIO

    doc = Document(BytesIO(conteudo))
    texto = _tabelas_texto(doc)

    assert code in texto
    assert quote.ordem_coleta.numero in texto
    assert quote.client_cnpj in texto  # CNPJ COLETA = CNPJ do cliente, nao da AMG
    assert quote.ordem_coleta.cnpj_filial not in texto  # CNPJ da AMG so escolhe o modelo/letterhead
    assert "Construtora Alfa" in texto  # nome da empresa na coleta
    assert "Obra Y" in texto  # destinatario, vindo da solicitacao
    assert "Rua das Flores, 100" in texto
    assert "R$ 11.520,00" in texto  # FE TOTAL EMPRESA = proposal.fe
    assert "5.000 kg" in texto  # peso: texto livre, repassado verbatim
    assert quote.data_coleta.strftime("%d/%m/%Y") in texto
    assert "Josemar Costa" in texto  # comercial responsavel, sempre fixo

    linhas = dict(_linhas_ordenadas(doc))
    # cotacao com 1 coleta/1 destino (o caso comum) -- rotulo do modelo
    # continua sem numero, como sempre foi (nao expande bloco nenhum).
    assert "ENDEREÇOS – LOCAL DA COLETA" in linhas
    assert "ENDEREÇOS – LOCAL DA COLETA 1" not in linhas
    assert "ENDEREÇOS – LOCAL DA ENTREGA" in linhas
    assert "ENDEREÇOS – LOCAL DA ENTREGA 1" not in linhas
    assert linhas["CNPJ ENTREGA"] == "98.765.432/0001-10"
    # FC SUGERIDO + PEDÁGIO: FC = 9600 (motorista 8000 + pedagio 900 +
    # impostos 700), exibido SEM o pedagio no 1o numero (8700) + o pedagio
    # separado (900) -- somados batem com o FC total, sem contar 2x. Esse
    # campo divide a linha "PESO" com ela (2o valor da linha, no ultimo
    # slot -- mesmo padrao de COMERCIAL RESPONSAVEL/N. DA COTACAO).
    peso_row = next(r for r in doc.tables[0].rows if r.cells[0].text.strip() == "PESO")
    assert peso_row.cells[-1].text.strip() == "R$ 8.700,00 + R$ 900,00"


def test_build_ordem_coleta_docx_fe_total_empresa_nao_inclui_adicionais(client):
    """FE TOTAL EMPRESA mostra proposal.fe -- nao proposal.valor_final (que
    soma adicionais tipo diaria/guincho). Os dois so coincidem quando nao ha
    adicional nenhum (caso do teste principal); este cobre o caso em que
    diferem de verdade."""
    from app.crud import get_quote_by_code
    from app.database import SessionLocal
    from app.ordem_coleta_docx import build_ordem_coleta_docx

    code = _fluxo_ate_oc(client, valor_diaria="450,00")  # FE=11520, valor_final=11970
    with SessionLocal() as db:
        quote = get_quote_by_code(db, code)
        assert quote.proposal.fe != quote.proposal.valor_final  # garante que o teste testa algo
        conteudo = build_ordem_coleta_docx(quote)

    from io import BytesIO

    doc = Document(BytesIO(conteudo))
    texto = _tabelas_texto(doc)
    assert "R$ 11.520,00" in texto  # FE, sem a diaria
    assert "R$ 11.970,00" not in texto  # valor_final nao aparece nesse campo


def test_build_ordem_coleta_docx_expande_blocos_de_coleta_e_entrega_extras(client):
    """O modelo .docx fisico so tem 1 bloco de coleta e 1 de entrega --
    com coletas/destinos extras, o gerador duplica o bloco (preservando
    formatacao, via clonagem do XML da linha) e numera os rotulos (so
    quando ha mais de 1 de cada lado), preenchendo cada copia com o
    endereco certo. Sem precisar de nenhum template novo por combinacao."""
    import app.crud as crud
    from app.database import SessionLocal
    from app.filiais import extrair_uf, sugestao_empresa_cnpj
    from app.ordem_coleta_docx import build_ordem_coleta_docx

    with SessionLocal() as db:
        quote = crud.create_quote(db, {
            "tipo_cotacao": "completa",
            "client_name": "Joao Alves", "client_company": "Construtora Alfa",
            "client_cnpj": "11.222.333/0001-44", "client_email": "compras@alfa.com.br",
            "client_phone": "(41) 99999-0000",
            "origem_cidade": "Curitiba - PR", "origem_cep": "80000-000",
            "origem_endereco": "Rua das Flores", "origem_numero": "100", "origem_bairro": "Centro",
            "destino_cidade": "Londrina - PR", "destino_cep": "86000-000",
            "destino_endereco": "Av. Brasil", "destino_numero": "200", "destino_bairro": "Jardim",
            "tipo_material": "Andaime", "descricao_material": None, "dimensoes": None,
            "qtd_volumes": 10, "peso_total": "5.000 kg", "valor_nf": None,
            "tipo_veiculo": "Carreta Sider", "carroceria": "Sider", "capacidade_aprox": "28 toneladas",
            "data_coleta": date.today() + timedelta(days=5), "data_entrega": None,
            "observacoes": None,
            "servico_diaria": True, "servico_guincho": False,
            "servico_ajudante": False, "servico_empilhadeira": False, "ajudante_qtd": None,
            "coletas_extras": [
                {"ordem": 2, "cidade": "Joinville - SC", "cep": "89200-000",
                 "endereco": "Rua das Palmeiras", "numero": "45", "bairro": "Centro"},
                {"ordem": 3, "cidade": "Blumenau - SC", "cep": None,
                 "endereco": None, "numero": None, "bairro": None},
            ],
            "destinos_extras": [
                {"ordem": 2, "cidade": "Maringá - PR", "cep": None,
                 "endereco": None, "numero": None, "bairro": None},
            ],
        })
        crud.create_or_update_solicitacao(db, quote, {
            "pagador": "Remetente", "pagador_documento": "12.345.678/0001-90",
            "fornecedor_nome": quote.client_company, "fornecedor_contato": None,
            "destinatario_nome": "Obra Central", "destinatario_contato": None,
            "valor_nf": None, "observacoes_operacionais": None,
        })
        uf = extrair_uf(quote.destino_cidade)
        empresa, cnpj = sugestao_empresa_cnpj(uf).split("|")
        crud.gerar_ordem_coleta(db, quote, empresa=empresa, cnpj_filial=cnpj, uf_referencia=uf, gerado_por="teste")
        db.refresh(quote)
        conteudo = build_ordem_coleta_docx(quote)

    from io import BytesIO

    doc = Document(BytesIO(conteudo))
    linhas = _linhas_ordenadas(doc)
    rotulos = [l for l, _ in linhas]

    # 3 blocos de coleta, 2 de entrega, numerados em sequencia
    assert "ENDEREÇOS – LOCAL DA COLETA 1" in rotulos
    assert "ENDEREÇOS – LOCAL DA COLETA 2" in rotulos
    assert "ENDEREÇOS – LOCAL DA COLETA 3" in rotulos
    assert "ENDEREÇOS – LOCAL DA COLETA 4" not in rotulos  # so 3, nao inventa um 4o
    assert "ENDEREÇOS – LOCAL DA ENTREGA 1" in rotulos
    assert "ENDEREÇOS – LOCAL DA ENTREGA 2" in rotulos
    assert "ENDEREÇOS – LOCAL DA ENTREGA 3" not in rotulos
    # sem a versao antiga (sem numero) sobrando
    assert "ENDEREÇOS – LOCAL DA COLETA" not in rotulos
    assert "ENDEREÇOS – LOCAL DA ENTREGA" not in rotulos

    def _endereco_apos(rotulo_bloco: str) -> str:
        i = rotulos.index(rotulo_bloco)
        # "ENDEREÇO COMPLETO..." e sempre a 3a linha depois do cabecalho do bloco
        for label, valor in linhas[i:i + 6]:
            if label.startswith("ENDEREÇO COMPLETO"):
                return valor
        raise AssertionError(f"endereco nao encontrado apos {rotulo_bloco!r}")

    assert _endereco_apos("ENDEREÇOS – LOCAL DA COLETA 1") == "Rua das Flores, 100, Centro, Curitiba - PR, CEP 80000-000"
    assert _endereco_apos("ENDEREÇOS – LOCAL DA COLETA 2") == "Rua das Palmeiras, 45, Centro, Joinville - SC, CEP 89200-000"
    assert _endereco_apos("ENDEREÇOS – LOCAL DA COLETA 3") == "Blumenau - SC"
    assert _endereco_apos("ENDEREÇOS – LOCAL DA ENTREGA 1") == "Av. Brasil, 200, Jardim, Londrina - PR, CEP 86000-000"
    assert _endereco_apos("ENDEREÇOS – LOCAL DA ENTREGA 2") == "Maringá - PR"

    # restante do documento (fora dos blocos expandidos) continua intacto
    texto = _tabelas_texto(doc)
    assert "DADOS DA CARGA" in texto
    assert "Carreta Sider" in texto
    assert conteudo[:2] == b"PK"


def test_build_ordem_coleta_docx_usa_modelo_da_empresa_escolhida(client):
    from app.crud import get_quote_by_code
    from app.database import SessionLocal
    from app.filiais import EMPRESA_TEMPLATE
    from app.ordem_coleta_docx import build_ordem_coleta_docx

    code = _fluxo_ate_oc(client)
    with SessionLocal() as db:
        quote = get_quote_by_code(db, code)
        empresa = quote.ordem_coleta.empresa
        conteudo = build_ordem_coleta_docx(quote)

    assert empresa in EMPRESA_TEMPLATE
    # o documento gerado deve ser um .docx (zip) valido, nao vazio
    assert conteudo[:2] == b"PK"
    assert len(conteudo) > 1000
